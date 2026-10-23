from docx import Document

doc = Document()
doc.add_heading("项目需求说明", 0)
doc.add_paragraph("本文档用于验证 TodayBlock 的 DOCX 文本提取与任务拆解。")
doc.add_heading("里程碑", level=1)
doc.add_paragraph("1. 10月28日完成需求评审，输出评审结论。")
doc.add_paragraph("2. 11月1日 18:00 前提交开发排期。")
doc.add_paragraph("3. 下周三前确认测试环境，联系运维开通权限。")
doc.add_paragraph("4. 11月8日交付第一版，安排客户验收。")
doc.save(r"D:\代码\today-block\samples\项目需求说明.docx")
print("docx written")
